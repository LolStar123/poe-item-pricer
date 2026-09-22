"""Extract the author's original spreadsheets without rewriting their formulas."""
import argparse
import json
import shutil
from pathlib import Path
from openpyxl import load_workbook

BOOKS=[('adorned','adorned profit.xlsx','Adorned roll valuation'),('meaning','light of meaning ID EV.xlsx','Light of Meaning identification'),('decks','stacked deck data.xlsx','Stacked-deck opening log')]

def export(source,catalogue):
    root=Path(__file__).resolve().parents[1]/'examples/portfolio';(root/'assets').mkdir(exist_ok=True);(root/'data').mkdir(exist_ok=True)
    books=[];cases=[]
    for key,name,title in BOOKS:
        path=source/name;formulas=load_workbook(path,read_only=True,data_only=False);values=load_workbook(path,read_only=True,data_only=True)
        sheets=[]
        for ws in formulas:
            cached=values[ws.title];cells=[]
            for row in ws:
                for cell in row:
                    if cell.value is not None:
                        cells.append({'address':cell.coordinate,'row':cell.row,'column':cell.column,'value':cached[cell.coordinate].value,'formula':cell.value if cell.data_type=='f' else None})
            sheets.append({'name':ws.title,'rows':max(c['row'] for c in cells),'columns':max(c['column'] for c in cells),'cells':cells})
        shutil.copyfile(path,root/'assets'/name)
        books.append({'id':key,'title':title,'file':'assets/'+name,'sheets':sheets})
        ws=values.active
        if key=='adorned':
            cases.append({'id':key,'title':title,'cost':ws['E5'].value,'unit':'original worksheet units','note':'Historical 50-150 roll model, with 101 equally weighted outcomes. Prices and zero valuations are archived author assumptions, not current trade quotes. The original sheet does not label the currency unit.','outcomes':[{'label':str(ws.cell(i,3).value)+'% roll','probability':1/101,'price':ws.cell(i,4).value,'source':'D'+str(i)} for i in range(5,106)]})
        if key=='meaning':
            cases.append({'id':key,'title':title,'cost':ws['G3'].value,'unit':'original worksheet units','note':'Thirteen equally weighted outcomes as assumed in the original sheet. Archived values are not current market quotes, and the source does not label the currency unit.','outcomes':[{'label':ws.cell(i,2).value,'probability':ws.cell(i,3).value,'price':ws.cell(i,4).value,'source':'D'+str(i)} for i in range(3,16)]})
        formulas.close();values.close()
    raw=json.loads(catalogue.read_text(encoding='utf-8'));mods=[{k:m[k] for k in ['mod_id','aura','display_text']} for m in raw['modifiers'] if m['role']=='aura']
    result={'books':books,'cases':cases,'modifiers':mods,'catalogueDate':raw['generated_at'],'league':raw['league']}
    (root/'data/archive.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(len(books),'original workbooks;',len(mods),'aura modifiers;',sum(len(c['outcomes']) for c in cases),'priced historical outcomes')

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('source',type=Path);p.add_argument('catalogue',type=Path);a=p.parse_args();export(a.source,a.catalogue)
